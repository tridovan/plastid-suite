"""The publication exports: multi-sheet Excel, the Supplementary package and the QC report.

Every one of them is regenerated from a run directory — the point of the phase — so the tests
build a run directory the way the CLI leaves one, and read the exports back.
"""

from __future__ import annotations

import importlib.util
import io
import json
from pathlib import Path

import pytest

from plastid.core.deps import ExportDependencyError
from plastid.report.package import build_package
from plastid.report.qc import render_qc, render_qc_for_run
from plastid.result.export import export_json
from plastid.result.figure import FIGURE_FORMATS
from plastid.result.figures import genome_page
from plastid.result.result import Column, Result, Scope, Table
from plastid.result.xlsx_export import SHEET_NAME_LIMIT, export_xlsx, sheet_name
from plastid.run.figures import render_run_figures

requires_excel = pytest.mark.skipif(
    importlib.util.find_spec("openpyxl") is None,
    reason="openpyxl is an optional dependency (pip install 'plastid-suite[excel]')",
)


def _genome_stats() -> Result:
    table = Table(
        id="genome_stats",
        title="Genome statistics",
        columns=(
            Column("organism", "Organism"),
            Column("accession", "Accession"),
            Column("length_bp", "Length", dtype="int", unit="bp"),
            Column("gc_pct", "GC", dtype="float", unit="%", precision=2),
            Column("n_genes_total", "Genes", dtype="int"),
            Column("n_cds", "CDS", dtype="int"),
            Column("n_trna", "tRNA", dtype="int"),
            Column("n_rrna", "rRNA", dtype="int"),
            Column("ir_length_bp", "IR", dtype="int", unit="bp"),
            Column("ir_identity_pct", "IR identity", dtype="float", unit="%", precision=2),
            Column("ir_source", "IR source"),
        ),
        rows=(
            {
                "organism": "Iris domestica",
                "accession": "NC_050833.1",
                "length_bp": 153_724,
                "gc_pct": 37.85,
                "n_genes_total": 132,
                "n_cds": 86,
                "n_trna": 38,
                "n_rrna": 8,
                "ir_length_bp": 26_214,
                "ir_identity_pct": 100.0,
                "ir_source": "annotation",
            },
        ),
        scope=Scope(label="whole plastome; both IR copies"),
    )
    return Result(module="genome_stats", module_version="0.1.0", tables={table.id: table})


def _structure() -> Result:
    regions = Table(
        id="regions",
        title="Regions",
        columns=(
            Column("region", "Region"),
            Column("start_1b", "Start", dtype="int", unit="bp"),
            Column("end_1b", "End", dtype="int", unit="bp"),
            Column("length_bp", "Length", dtype="int", unit="bp"),
            Column("gc_pct", "GC", dtype="float", unit="%", precision=2),
        ),
        rows=(
            {"region": "LSC", "start_1b": 1, "end_1b": 83_127, "length_bp": 83_127, "gc_pct": 36.0},
            {
                "region": "IRa",
                "start_1b": 83_128,
                "end_1b": 109_341,
                "length_bp": 26_214,
                "gc_pct": 43.0,
            },
        ),
        scope=Scope(label="normalized frame", coord_system="normalized"),
        key="region",
    )
    summary = Table(
        id="structure_summary",
        title="Structure",
        columns=(
            Column("accession", "Accession"),
            Column("rotation_offset", "Rotation offset", dtype="int", unit="bp"),
        ),
        rows=({"accession": "NC_050833.1", "rotation_offset": 0},),
        scope=Scope(label="normalized frame", coord_system="normalized"),
    )
    skew = Table(
        id="gc_skew_windows",
        title="GC skew",
        columns=(
            Column("region", "Region"),
            Column("window_start_1b", "Start", dtype="int", unit="bp"),
            Column("window_end_1b", "End", dtype="int", unit="bp"),
            Column("gc_skew", "GC skew", dtype="float", precision=4),
        ),
        rows=(
            {"region": "LSC", "window_start_1b": 1, "window_end_1b": 1000, "gc_skew": -0.05},
            {"region": "LSC", "window_start_1b": 1001, "window_end_1b": 2000, "gc_skew": None},
        ),
        scope=Scope(label="normalized frame", coord_system="normalized"),
    )
    return Result(
        module="structure",
        module_version="0.2.0",
        tables={"regions": regions, "structure_summary": summary, "gc_skew_windows": skew},
    )


def _run_directory(root: Path) -> Path:
    """A run directory shaped the way the CLI leaves one: one module dir per module."""
    run = root / "standin" / "20261005-000000-abcdef12"
    for order, result in enumerate((_genome_stats(), _structure()), start=1):
        module_dir = run / f"{order:02d}_{result.module}"
        module_dir.mkdir(parents=True)
        (module_dir / "result.json").write_text(export_json(result), encoding="utf-8")
    (run / "manifest.json").write_text(
        json.dumps(
            {
                "run_id": run.name,
                "command": ["plastid", "run", "--module", "genome_stats"],
                "package_version": "0.1.0",
                "started_at": "2026-10-05T00:00:00",
                "finished_at": "2026-10-05T00:00:01",
                "status": "ok",
                "params": {"window": 1000},
                "inputs": [
                    {"path": "in.gb", "sha256": "a" * 64, "bytes": 12},
                ],
                "coordinate_system": {"name": "normalized", "rotation_offset": 0},
                "ir": {
                    "gate_passed": True,
                    "length": 26_214,
                    "identity": 1.0,
                    "source": "annotation",
                },
                "modules": [
                    {
                        "id": "genome_stats",
                        "version": "0.1.0",
                        "status": "ok",
                        "files": [],
                        "tools": [
                            {
                                "name": "mafft",
                                "version": "7.526",
                                "path": "/usr/bin/mafft",
                                "exit_code": 0,
                            }
                        ],
                    }
                ],
                "files": [
                    {
                        "path": "01_genome_stats/result.json",
                        "module": "genome_stats",
                        "sha256": "b" * 64,
                        "bytes": 10,
                    }
                ],
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    return run


# --------------------------------------------------------------------------- Excel


def test_a_sheet_name_is_sanitized_truncated_and_unique() -> None:
    assert sheet_name("ssr", "ssrs") == "ssr.ssrs"
    assert "[" not in sheet_name("a", "b[c]")
    long = sheet_name("module", "x" * 40)
    assert len(long) <= SHEET_NAME_LIMIT and long.endswith("0") is False
    taken = {long}
    other = sheet_name("module", "x" * 40, taken)
    assert other != long and len(other) <= SHEET_NAME_LIMIT


@requires_excel
def test_the_workbook_holds_one_sheet_per_table_and_is_byte_identical() -> None:
    results = {"genome_stats": _genome_stats(), "structure": _structure()}
    first = export_xlsx(results)
    assert first == export_xlsx(results), "the workbook must regenerate byte-identically"

    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(first))
    assert book.sheetnames == [
        "genome_stats.genome_stats",
        "structure.regions",
        "structure.structure_summary",
        "structure.gc_skew_windows",
    ]
    sheet = book["genome_stats.genome_stats"]
    assert [cell.value for cell in sheet[1]][:3] == ["Organism", "Accession", "Length (bp)"]
    assert sheet["C2"].value == "153724"  # the cell's own rendering, not a second rounding
    assert book.properties.created.year == 1980


def test_the_workbook_pins_both_timestamps() -> None:
    """`modified` too — and it is the one openpyxl overwrites.

    Measured 10/10/2026 (§9.C-38): `ExcelWriter.write_workbook` sets
    ``workbook.properties.modified = datetime.datetime.now(...)`` **after** the caller has pinned
    it, so the pin on that field never reached the file and the byte-identity test above failed
    once in a full run (3 times in 300 exports — a clock-second race). This asserts the contract
    directly instead of hoping the race shows up.
    """
    import zipfile

    assert importlib.util.find_spec("openpyxl") is not None
    archive = zipfile.ZipFile(io.BytesIO(export_xlsx({"genome_stats": _genome_stats()})))
    core = archive.read("docProps/core.xml").decode("utf-8")
    assert core.count("1980-01-01T00:00:00Z") == 2, (
        f"created and modified must both carry the pinned stamp, got: {core}"
    )


def test_a_core_xml_without_the_stamp_stops_the_export() -> None:
    """The negative case: openpyxl changing the shape of ``core.xml`` must be loud (§9.C-38).

    Silently returning the file unchanged is the failure this helper removes — a workbook whose
    bytes depend on the second it was written in, promised as reproducible.
    """
    from plastid.result.xlsx_export import _pin_modified

    with pytest.raises(ValueError, match="dcterms:modified"):
        _pin_modified(b"<cp:coreProperties><dcterms:created>1980</dcterms:created></cp:coreProperties>")


def test_a_missing_excel_extra_is_a_named_error(monkeypatch: pytest.MonkeyPatch) -> None:
    import sys

    monkeypatch.setitem(sys.modules, "openpyxl", None)
    with pytest.raises(ExportDependencyError, match="plastid-suite\\[excel\\]"):
        export_xlsx({"genome_stats": _genome_stats()})


# ------------------------------------------------------------------------- package


def test_the_package_regenerates_from_a_run_directory_alone(tmp_path: Path) -> None:
    pytest.importorskip("openpyxl", reason="the 'excel' extra is not installed")
    run = _run_directory(tmp_path)
    package = build_package(
        run, results={"genome_stats": _genome_stats(), "structure": _structure()}
    )
    names = sorted(path.relative_to(package.directory).as_posix() for path in package.files)
    assert "tables.xlsx" in names
    assert "index.md" in names and "MANIFEST.json" in names
    assert "tables/genome_stats__genome_stats.csv" in names
    assert len([name for name in names if name.startswith("tables/")]) == 4

    manifest = json.loads((package.directory / "MANIFEST.json").read_text(encoding="utf-8"))
    # Every file but the manifest itself: a file cannot carry its own digest.
    listed = {entry["path"] for entry in manifest["files"]}
    written = {path.relative_to(package.directory).as_posix() for path in package.files}
    assert listed == written - {"MANIFEST.json"}
    index = (package.directory / "index.md").read_text(encoding="utf-8")
    assert "stats_summary" not in index  # only tables this run holds
    assert "| genome_stats | genome_stats |" in index
    assert "no figure's source tables are in this run" in index


def test_the_package_needs_no_plotting_library_to_carry_tables(tmp_path: Path) -> None:
    """The package's tables are text and a workbook; a minimum install can still produce them."""
    run = _run_directory(tmp_path)
    package = build_package(run, results={"genome_stats": _genome_stats()})
    assert (package.directory / "tables" / "genome_stats__genome_stats.csv").is_file()


# ------------------------------------------------------------------------- QC report


def test_the_qc_report_names_every_tool_version_and_input_hash(tmp_path: Path) -> None:
    run = _run_directory(tmp_path)
    manifest = json.loads((run / "manifest.json").read_text(encoding="utf-8"))
    text = render_qc(
        manifest,
        figures=("genome_page.pdf",),
        skipped=(("pi_window", "needs diversity.diversity_windows — not in this run"),),
        audits=(("genome_stats", True, 7),),
    )
    assert "mafft" in text and "7.526" in text and "/usr/bin/mafft" in text
    assert "a" * 64 in text  # the input's sha256
    assert "genome_page.pdf" in text
    assert "skipped `pi_window`" in text
    assert "Queryable-number audit" in text and "clean" in text
    assert "rotation offset 0" in text


def test_the_qc_report_is_written_into_the_run_directory(tmp_path: Path) -> None:
    run = _run_directory(tmp_path)
    text, target = render_qc_for_run(run, figures=("genome_page.pdf",))
    assert target == run / "report" / "qc.md"
    assert target.read_text(encoding="utf-8") == text
    assert "# Technical QC report" in text


# --------------------------------------------------------------------- genome page


def test_the_page_series_equals_the_cells_and_prints_no_number_of_its_own() -> None:
    results = {"genome_stats": _genome_stats(), "structure": _structure()}
    data = genome_page.series(results)
    summary = results["genome_stats"].tables["genome_stats"]
    assert data.length == summary.format("length_bp")
    assert data.gc == summary.format("gc_pct")
    for line in data.stat_lines():
        for piece in (data.length, data.gc, data.genes, data.cds, data.trna, data.rrna):
            assert piece in line or piece not in line  # the line is built from these strings
    assert data.stat_lines()[0] == "Length 153724 bp · GC 37.85 %"
    assert data.regions[0].length == results["structure"].tables["regions"].format("LSC.length_bp")
    assert [point.gc_skew for point in data.skew] == [-0.05, None]


def test_the_page_is_pdf_only_and_the_others_carry_all_three_formats() -> None:
    from plastid.result.figures import get

    assert genome_page.SPEC.formats == ("pdf",)
    assert get("circular_map").formats == FIGURE_FORMATS


def test_a_pdf_only_figure_is_skipped_when_only_svg_was_asked_for(tmp_path: Path) -> None:
    pytest.importorskip("matplotlib", reason="the 'figures' extra is not installed")
    run = _run_directory(tmp_path)
    outcome = render_run_figures(run, formats=("svg",))
    skipped = dict(outcome.skipped)
    assert "genome_page" in skipped and "pdf only" in skipped["genome_page"]


def test_the_run_level_pass_writes_the_page_as_a_pdf(tmp_path: Path) -> None:
    pytest.importorskip("matplotlib", reason="the 'figures' extra is not installed")
    run = _run_directory(tmp_path)
    outcome = render_run_figures(run)
    names = sorted(path.name for path in outcome.files)
    assert "genome_page.pdf" in names
    assert not any(name.startswith("genome_page.") and not name.endswith(".pdf") for name in names)
    assert (run / "figures" / "genome_page.pdf").stat().st_size > 5_000
    skipped = dict(outcome.skipped)
    assert "pi_window" in skipped  # a figure with no source tables is still named
