"""Word export: reproducibility, package health, and the layout decisions that carry meaning."""

from __future__ import annotations

import importlib.util
import io
import sys
import zipfile
from dataclasses import replace

import pytest

from plastid.core.deps import ExportDependencyError
from plastid.result.docx_export import DocxMeta, export_docx
from plastid.result.result import Column, Result, Scope, Table

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

requires_docx = pytest.mark.skipif(
    importlib.util.find_spec("docx") is None,
    reason="the 'docx' extra is not installed — pip install 'plastid-suite[docx]'",
)


def _table(rows: int = 2) -> Table:
    return Table(
        id="summary",
        title="Summary",
        columns=(
            Column("name", "Name"),
            Column("length_bp", "Length", dtype="int", unit="bp"),
        ),
        rows=tuple({"name": f"row{index}", "length_bp": 100 + index} for index in range(rows)),
        scope=Scope(label="whole plastome", coord_system="genbank"),
        key="name" if rows > 1 else None,
    )


def _result(table: Table | None = None) -> Result:
    return Result(
        module="genome_stats",
        module_version="0.1.0",
        tables={(table or _table()).id: table or _table()},
        notes=("A note the document must carry.",),
    )


def _meta() -> DocxMeta:
    return DocxMeta(
        title="Genome statistics",
        label="NC_050833.1 · genome_stats",
        prose="## Heading\n\nA **bold** claim about `genes` and an *italic* one.\n\n- first\n",
        provenance=(("Dataset", "golden"), ("Module", "genome_stats 0.1.0")),
        references=("Biopython — Cock et al. (2009).",),
    )


def _parts(blob: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


# ------------------------------------------------------------------- reproducibility


@requires_docx
def test_two_builds_of_the_same_result_are_byte_identical() -> None:
    """P2-09's Done-when, stated as a test: regenerate the document, get the same file."""
    assert export_docx(_result(), _meta()) == export_docx(_result(), _meta())


@requires_docx
def test_every_package_entry_carries_a_fixed_timestamp() -> None:
    """The mechanism behind the test above — python-docx stamps the wall clock, we do not."""
    with zipfile.ZipFile(io.BytesIO(export_docx(_result(), _meta()))) as archive:
        entries = archive.infolist()
    assert [entry.date_time for entry in entries] == [(1980, 1, 1, 0, 0, 0)] * len(entries)
    assert [entry.filename for entry in entries] == sorted(entry.filename for entry in entries)
    assert {entry.create_system for entry in entries} == {0}


@requires_docx
def test_a_different_result_produces_a_different_document() -> None:
    """The determinism above must not be the determinism of an empty document."""
    first = export_docx(_result(_table(rows=2)), _meta())
    second = export_docx(_result(_table(rows=3)), _meta())
    assert first != second


# ------------------------------------------------------------------------- package health


@requires_docx
def test_the_pruned_package_has_no_dangling_references() -> None:
    """Pruning template parts is only safe if nothing still points at them."""
    from lxml import etree

    parts = _parts(export_docx(_result(), _meta()))
    assert not [name for name in parts if name.startswith("customXml/")]
    assert "docProps/thumbnail.jpeg" not in parts
    assert "word/stylesWithEffects.xml" not in parts

    for name, blob in parts.items():
        if not name.endswith(".rels"):
            continue
        base = name.rsplit("_rels/", 1)[0].rstrip("/")
        for relationship in etree.fromstring(blob):
            if relationship.get("TargetMode") == "External":
                continue
            target = relationship.get("Target", "")
            resolved = f"{base}/{target}" if base else target
            assert resolved.lstrip("/") in parts, f"{name} points at missing {target}"


@requires_docx
def test_every_style_the_document_uses_survived_the_prune() -> None:
    """A missing style is invisible until Word renders the text wrong; assert it structurally."""
    from lxml import etree

    parts = _parts(export_docx(_result(), _meta()))
    defined = set(_defined_styles(parts["word/styles.xml"]))
    used = {
        node.get(f"{W}val")
        for node in etree.fromstring(parts["word/document.xml"]).iter(
            f"{W}pStyle", f"{W}rStyle", f"{W}tblStyle"
        )
    }
    assert used - {None} <= defined
    assert {"Normal", "Title", "Heading1", "ListBullet"} <= defined


@requires_docx
def test_pruning_keeps_a_handful_of_styles_out_of_the_templates_hundreds() -> None:
    """The template ships Word's whole style catalogue; the deliverable should not carry it."""

    template_styles = len(_template_styles())
    shipped = len(_defined_styles(_parts(export_docx(_result(), _meta()))["word/styles.xml"]))
    assert template_styles > 100  # if a future template ships fewer, this is not the guard
    assert shipped < 25, f"{shipped} styles kept out of {template_styles} — is the prune working?"


def _template_styles() -> list[str]:
    from docx import Document

    buffer = io.BytesIO()
    Document().save(buffer)
    return _defined_styles(_parts(buffer.getvalue())["word/styles.xml"])


def _defined_styles(blob: bytes) -> list[str]:
    from lxml import etree

    return [style.get(f"{W}styleId") for style in etree.fromstring(blob).findall(f"{W}style")]


@requires_docx
def test_the_written_document_reopens_and_keeps_its_numbers() -> None:
    from docx import Document

    document = Document(io.BytesIO(export_docx(_result(), _meta())))
    text = [paragraph.text for paragraph in document.paragraphs]
    text += [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    assert "Genome statistics" in text
    assert "row1" in text and "101" in text
    assert "A note the document must carry." in text
    assert "Biopython — Cock et al. (2009)." in text
    assert "Dataset" in text and "golden" in text


@requires_docx
def test_a_changed_cell_changes_the_document() -> None:
    """Numbers reach the page from cells, never from a literal in the exporter."""
    table = _table()
    changed = replace(table, rows=({"name": "row0", "length_bp": 999},))
    assert "101" in _text(export_docx(_result(table), _meta()))
    assert "999" not in _text(export_docx(_result(table), _meta()))
    assert "999" in _text(export_docx(_result(changed), _meta()))


def _text(blob: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(blob))
    parts = [paragraph.text for paragraph in document.paragraphs]
    parts += [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join(parts)


# ------------------------------------------------------------------------------- layout


@requires_docx
def test_wide_tables_get_their_own_landscape_section() -> None:
    from docx import Document
    from docx.enum.section import WD_ORIENT

    wide = Table(
        id="wide",
        title="Wide",
        columns=tuple(Column(f"c{index}", f"Column number {index}") for index in range(10)),
        rows=tuple({f"c{index}": "value" for index in range(10)} for _ in range(40)),
        scope=Scope(label="whole plastome"),
    )
    document = Document(io.BytesIO(export_docx(_result(wide), _meta())))
    orientations = [section.orientation for section in document.sections]
    assert orientations[0] == WD_ORIENT.PORTRAIT
    assert WD_ORIENT.LANDSCAPE in orientations


@requires_docx
def test_the_header_row_repeats_and_rows_do_not_split() -> None:
    from docx import Document
    from docx.oxml.ns import qn

    document = Document(io.BytesIO(export_docx(_result(_table(rows=40)), _meta())))
    grid = next(t for t in document.tables if t.rows[0].cells[0].text == "Name")
    rows = list(grid._tbl.iter(qn("w:tr")))
    repeat = [row.find(qn("w:trPr")).find(qn("w:tblHeader")) is not None for row in rows]
    split = [row.find(qn("w:trPr")).find(qn("w:cantSplit")) is not None for row in rows]
    assert repeat[0] is True and not any(repeat[1:]), "only the header repeats"
    assert all(split), "a row broken across a page boundary reads as two rows"


@requires_docx
def test_the_emphasised_row_is_bold() -> None:
    from docx import Document

    table = replace(
        _table(rows=3),
        rows=_table(rows=3).rows + ({"name": "Total", "length_bp": 300},),
        emphasis=("Total",),
    )
    document = Document(io.BytesIO(export_docx(_result(table), _meta())))
    word_table = next(t for t in document.tables if t.rows[0].cells[0].text == "Name")
    bold_by_row = [
        {run.bold for paragraph in cell.paragraphs for run in paragraph.runs}
        for cell in word_table.rows[-1].cells
    ]
    assert bold_by_row == [{True}, {True}]


# --------------------------------------------------------------------- optional dependency


def test_a_missing_optional_dependency_is_a_named_error(monkeypatch: pytest.MonkeyPatch) -> None:
    """The minimum install must stay pure Python — and say so, not raise ImportError."""
    monkeypatch.setitem(sys.modules, "docx", None)
    with pytest.raises(ExportDependencyError, match="plastid-suite\\[docx\\]"):
        export_docx(_result(), _meta())


# ------------------------------------------------------------------------- markdown subset


@requires_docx
def test_the_markdown_subset_renders_as_word_blocks() -> None:
    from docx import Document

    meta = replace(
        _meta(),
        prose=(
            "## Section\n\n"
            "A **bold** claim and an *italic* one and `code`.\n"
            "A hard wrap must continue the same paragraph.\n\n"
            "- first\n- second\n"
        ),
    )
    document = Document(io.BytesIO(export_docx(_result(), meta)))
    paragraphs = document.paragraphs
    headings = [p.text for p in paragraphs if p.style.name.startswith("Heading")]
    assert "Section" in headings
    body = next(p for p in paragraphs if p.text.startswith("A bold claim"))
    assert body.text.endswith("continue the same paragraph.")
    assert {run.bold for run in body.runs if run.text == "bold"} == {True}
    assert {run.italic for run in body.runs if run.text == "italic"} == {True}
    assert [p.text for p in paragraphs if p.style.name == "List Bullet"][:2] == [
        "first",
        "second",
    ]


@requires_docx
def test_a_pipe_table_in_the_prose_becomes_a_word_table() -> None:
    from docx import Document

    meta = replace(_meta(), prose="Intro\n\n| A | B |\n|---|---|\n| 1 | 2 |\n")
    document = Document(io.BytesIO(export_docx(_result(), meta)))
    tables = [t for t in document.tables if t.rows[0].cells[0].text == "A"]
    assert len(tables) == 1
    assert tables[0].rows[1].cells[1].text == "2"


# ------------------------------------------------------------------------------- helpers


@requires_docx
def test_the_document_directory_has_no_leftover_default_style_bloat() -> None:
    """A guard against a future python-docx template that is larger again: state the budget."""
    blob = export_docx(_result(), _meta())
    assert len(blob) < 60_000


def test_importing_the_package_never_pulls_in_the_optional_extra() -> None:
    """The minimum install is biopython + PyYAML. Checked at runtime, in a clean interpreter.

    A textual scan of the source would miss `import docx` hidden behind an indirection; running
    the import in a fresh process and looking at `sys.modules` is the claim itself.
    """
    import subprocess

    program = (
        "import sys, plastid.cli, plastid.modules, plastid.result.docx_export;"
        " heavy = sorted(m for m in ('docx', 'lxml') if m in sys.modules);"
        " print(','.join(heavy))"
    )
    result = subprocess.run(
        [sys.executable, "-c", program], capture_output=True, text=True, check=True
    )
    assert result.stdout.strip() == "", f"importing plastid pulled in {result.stdout.strip()}"


# --------------------------------------------------------------------- joining documents


def _document(title: str, *, rows: int = 2, wide: bool = False) -> bytes:
    columns = (Column("name", "Name"), Column("length_bp", "Length", dtype="int", unit="bp"))
    if wide:
        columns = tuple(Column(f"c{i}", f"Column number {i}") for i in range(14))
    table = Table(
        id="t",
        title="A table",
        columns=columns,
        rows=tuple(
            {column.name: f"{title}-{index}" for column in columns} for index in range(rows)
        ),
        scope=Scope(label="whole plastome"),
        key=columns[0].name,
    )
    result = Result(module="m", module_version="0.1.0", tables={"t": table})
    return export_docx(result, replace(_meta(), title=title))


@requires_docx
def test_merging_keeps_every_document_and_is_reproducible() -> None:
    from plastid.result.docx_export import merge_documents

    blob = merge_documents([_document("First"), _document("Second"), _document("Third")])
    assert blob == merge_documents([_document("First"), _document("Second"), _document("Third")])
    text = _text(blob)
    assert "First" in text and "Second" in text and "Third" in text
    assert "First-0" in text and "Second-0" in text and "Third-0" in text


@requires_docx
def test_merging_one_document_returns_it_untouched() -> None:
    from plastid.result.docx_export import merge_documents

    single = _document("Only")
    assert merge_documents([single]) == single


@requires_docx
def test_merging_keeps_each_documents_own_page_setup() -> None:
    """Dropping the earlier documents' sectPr would re-flow a landscape table into portrait."""
    from docx import Document
    from docx.enum.section import WD_ORIENT

    from plastid.result.docx_export import merge_documents

    merged = merge_documents([_document("Portrait"), _document("Wide", wide=True)])
    orientations = [section.orientation for section in Document(io.BytesIO(merged)).sections]
    assert WD_ORIENT.PORTRAIT in orientations and WD_ORIENT.LANDSCAPE in orientations
    assert orientations[-1] == WD_ORIENT.LANDSCAPE


@requires_docx
def test_merging_sets_the_document_title() -> None:
    from docx import Document

    from plastid.result.docx_export import merge_documents

    merged = merge_documents([_document("A"), _document("B")], title="Run report")
    assert Document(io.BytesIO(merged)).core_properties.title == "Run report"


def test_merging_rejects_a_document_this_exporter_did_not_build() -> None:
    from plastid.result.docx_export import merge_documents

    foreign = io.BytesIO()
    with zipfile.ZipFile(foreign, "w") as archive:
        archive.writestr("word/document.xml", "<w:document/>")
    with pytest.raises(ValueError, match="different part list"):
        merge_documents([_document("Built"), foreign.getvalue()])
    with pytest.raises(ValueError, match="nothing to merge"):
        merge_documents([])


@requires_docx
def test_merging_relabels_the_footer() -> None:
    """The merged document keeps the first module's footer, which would name the wrong module."""
    from docx import Document

    from plastid.result.docx_export import merge_documents

    merged = merge_documents([_document("A"), _document("B")], label="the whole run")
    footers = [
        paragraph.text
        for section in Document(io.BytesIO(merged)).sections
        for paragraph in section.footer.paragraphs
    ]
    assert any("the whole run" in text for text in footers), footers
    assert not any("genome_stats ·" in text for text in footers)
