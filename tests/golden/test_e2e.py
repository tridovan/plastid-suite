"""End-to-end golden run: one command produces a complete, traceable run directory.

Covers P1-08's Done-when: run dir with 00_run_info + tables + report; tool/parameter report;
cache hit on an identical second run; byte-identical restore after deleting an output.
"""

from __future__ import annotations

import importlib
import importlib.util
import json
from pathlib import Path

import pytest
from Bio import SeqIO

from conftest import require_tools
from plastid.cli import main
from plastid.run.manifest import validate


def _run(
    tmp_path: Path,
    source: Path,
    dataset: str = "iris_domestica",
    extra: list[str] | None = None,
    module: str | None = "genome_stats",
) -> int:
    """``module=None`` leaves the module flags to ``extra`` — for multi-module runs."""
    return main(
        [
            "run",
            *(["--module", module] if module else []),
            "--input",
            str(source),
            "--out",
            str(tmp_path),
            "--dataset",
            dataset,
            *(extra or []),
        ]
    )


def _run_dirs(tmp_path: Path, dataset: str = "iris_domestica") -> set[Path]:
    """Every run directory the dataset currently holds.

    A run id starts with a **wall-clock** stamp, so sorting the directories by name is not
    sorting them by time: measured 08/10/2026, this machine's clock stepped ~1 s backwards
    between two runs of this very test, and `sorted(...)[2]` then pointed at the second run —
    whose output the test had just deleted. Identify a run's directory by
    the difference before/after it, never by its name.
    """
    return set((tmp_path / dataset).iterdir())


def _only_run_dir(tmp_path: Path, dataset: str = "iris_domestica") -> Path:
    runs = _run_dirs(tmp_path, dataset)
    assert len(runs) == 1
    return runs.pop()


def test_run_produces_a_complete_traceable_directory(tmp_path: Path, data_dir: Path) -> None:
    assert _run(tmp_path, data_dir / "NC_050833.1.gb") == 0
    run_dir = _only_run_dir(tmp_path)

    manifest = json.loads((run_dir / "manifest.json").read_text(encoding="utf-8"))
    validate(manifest)
    assert manifest["status"] == "ok"
    assert manifest["ir"]["gate_passed"] is True
    assert manifest["ir"]["length"] == 26_214
    assert manifest["coordinate_system"]["rotation_offset"] == 0
    assert manifest["modules"][0]["id"] == "genome_stats"
    assert manifest["files"], "the manifest must index every produced file"
    for entry in manifest["files"]:
        assert (run_dir / entry["path"]).is_file()

    stats = (run_dir / "01_genome_stats" / "tables" / "genome_stats.csv").read_text(
        encoding="utf-8"
    )
    assert "153724" in stats
    assert "26214" in stats
    assert "# scope:" in stats  # every export carries its scope

    report = (run_dir / "01_genome_stats" / "report.md").read_text(encoding="utf-8")
    assert "Draft — needs human review." in report
    assert "132 genes" in report
    assert "### References" in report  # biopython citation resolved from the registry

    run_info = (run_dir / "00_run_info" / "run_info.md").read_text(encoding="utf-8")
    assert "## Parameters" in run_info
    assert "## IR gate" in run_info
    assert "passed" in run_info
    assert "genome_stats" in run_info


def test_second_identical_run_is_cached_and_identical(tmp_path: Path, data_dir: Path) -> None:
    source = data_dir / "NC_050833.1.gb"
    assert _run(tmp_path, source) == 0
    first_dir = _only_run_dir(tmp_path)
    first_csv = (first_dir / "01_genome_stats" / "tables" / "genome_stats.csv").read_bytes()

    before = _run_dirs(tmp_path)
    assert _run(tmp_path, source) == 0
    runs = _run_dirs(tmp_path)
    assert len(runs) == 2, "each run gets its own directory"
    second_dir = (runs - before).pop()
    manifest = json.loads((second_dir / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["modules"][0]["status"] == "cached"

    second_csv_path = second_dir / "01_genome_stats" / "tables" / "genome_stats.csv"
    assert second_csv_path.read_bytes() == first_csv  # same input, same numbers

    second_csv_path.unlink()
    before = _run_dirs(tmp_path)
    assert _run(tmp_path, source) == 0
    third_dir = (_run_dirs(tmp_path) - before).pop()
    restored = (third_dir / "01_genome_stats" / "tables" / "genome_stats.csv").read_bytes()
    assert restored == first_csv  # restored byte-identically from the cache


def test_failed_gate_stops_the_run_with_exit_2(tmp_path: Path, data_dir: Path) -> None:
    # NC_063921's IR copies are 99.16 % identical end to end — position-based analysis must not run.
    records = SeqIO.parse(str(data_dir / "iridaceae_edge_cases.gb"), "genbank")
    target = next(record for record in records if record.id == "NC_063921.1")
    single = tmp_path / "NC_063921.1.gb"
    SeqIO.write(target, str(single), "genbank")

    assert _run(tmp_path, single, dataset="diverged") == 2
    run_dir = _only_run_dir(tmp_path, "diverged")
    manifest = json.loads((run_dir / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["status"] == "failed"
    assert not (run_dir / "01_genome_stats").exists(), "no module output after a failed gate"


def test_golden_numbers_reach_the_report_and_the_csv(tmp_path: Path, data_dir: Path) -> None:
    """P2-01's Done-when: length · 132 = 86+38+8 · GC, in both the report and the CSV."""
    assert _run(tmp_path, data_dir / "NC_050833.1.gb") == 0
    run_dir = _only_run_dir(tmp_path)

    report = " ".join((run_dir / "01_genome_stats" / "report.md").read_text("utf-8").split())
    assert "153724 bp" in report
    assert "37.85 %" in report
    assert "132 genes: 86 protein-coding, 38 tRNA and 8 rRNA genes." in report

    lines = (run_dir / "01_genome_stats" / "tables" / "genome_stats.csv").read_text("utf-8")
    assert "# scope:" in lines  # every exported number carries its scope label
    header, row = [line for line in lines.splitlines() if not line.startswith("#")][:2]
    assert header.split(",")[2:8] == ["Length (bp)", "GC (%)", "Genes", "CDS", "tRNA", "rRNA"]
    assert row.split(",")[2:8] == ["153724", "37.85", "132", "86", "38", "8"]


def test_second_golden_genome_runs_cleanly(tmp_path: Path, data_dir: Path) -> None:
    """The comparison genome NC_024936.1 runs end-to-end on the explicit sequence path.

    Its annotated copies are 26,220/26,186 bp and align at 99.80 % end to end — below the gate, so
    the default annotation-first run refuses it (D11-2). The sequence's own repeat is used only when
    the user asks for it; this is that labelled path.
    """
    assert (
        _run(
            tmp_path,
            data_dir / "NC_024936.1.gb",
            dataset="iris_gatesii",
            extra=["--set", "ir_source=sequence"],
        )
        == 0
    )
    run_dir = _only_run_dir(tmp_path, "iris_gatesii")
    manifest = json.loads((run_dir / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["status"] == "ok"
    assert manifest["params"]["ir_source"] == "sequence"
    assert manifest["ir"]["gate_passed"] is True
    assert manifest["ir"]["source"] == "denovo_diag"
    # the annotation it did not use is still recorded, so the choice is visible
    assert manifest["ir"]["annotated_lengths"] == [26_220, 26_186]
    assert manifest["ir"]["alignment_method"] == "anchored"
    assert manifest["ir"]["known_columns"] >= manifest["ir"]["compared_columns"] // 2


def test_the_annotation_source_is_refused_with_the_way_out(
    tmp_path: Path, data_dir: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    """The default path refuses NC_024936.1 and names ``ir_source=sequence`` (D11-2)."""
    assert _run(tmp_path, data_dir / "NC_024936.1.gb", dataset="iris_gatesii_refused") == 2
    stderr = capsys.readouterr().err
    assert "99.80 % identity" in stderr
    assert "--set ir_source=sequence" in stderr
    run_dir = _only_run_dir(tmp_path, "iris_gatesii_refused")
    manifest = json.loads((run_dir / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["status"] == "failed"
    assert manifest["ir"]["gate_passed"] is False
    assert manifest["ir"]["annotated_lengths"] == [26_220, 26_186]


def test_only_the_diverged_ir_genome_stops_after_gap_aware_measurement(
    tmp_path: Path, data_dir: Path
) -> None:
    """NC_063920 now passes; NC_063921 still fails. Both verdicts come from the full copies.

    Re-measured 26/10/2026 with the end-to-end alignment (M11.4): NC_063920's copies differ by 10
    mismatches and 5 internal gaps over 26,164 columns (99.94 %), while NC_063921's differ by 147
    mismatches and 72 gaps over 26,219 (99.16 %). The old 92.63 % for NC_063920 was the ungapped
    mirror pairing shifting every base after a 1 bp indel.
    """
    records = {
        record.id: record
        for record in SeqIO.parse(str(data_dir / "iridaceae_edge_cases.gb"), "genbank")
        if record.id in {"NC_063920.1", "NC_063921.1"}
    }
    assert set(records) == {"NC_063920.1", "NC_063921.1"}

    passing = tmp_path / "NC_063920.1.gb"
    SeqIO.write(records["NC_063920.1"], str(passing), "genbank")
    assert _run(tmp_path, passing, dataset="nc_063920") == 0
    passed_manifest = json.loads(
        (_only_run_dir(tmp_path, "nc_063920") / "manifest.json").read_text(encoding="utf-8")
    )
    assert passed_manifest["ir"]["gate_passed"] is True
    assert passed_manifest["ir"]["identity"] == pytest.approx(26_149 / 26_164)

    failing = tmp_path / "NC_063921.1.gb"
    SeqIO.write(records["NC_063921.1"], str(failing), "genbank")
    assert _run(tmp_path, failing, dataset="nc_063921") == 2
    failed_manifest = json.loads(
        (_only_run_dir(tmp_path, "nc_063921") / "manifest.json").read_text(encoding="utf-8")
    )
    assert failed_manifest["status"] == "failed"
    assert not (_only_run_dir(tmp_path, "nc_063921") / "01_genome_stats").exists()


def test_structure_run_reports_regions_and_skew(tmp_path: Path, data_dir: Path) -> None:
    """P2-02 end-to-end: the region table, its scope label, and the skew convention."""
    assert _run(tmp_path, data_dir / "NC_050833.1.gb", module="structure") == 0
    run_dir = _only_run_dir(tmp_path)

    lines = (run_dir / "01_structure" / "tables" / "regions.csv").read_text("utf-8")
    assert "coordinates: normalized" in lines
    body = dict(
        (fields[0], fields[1:])
        for fields in (line.split(",") for line in lines.splitlines() if not line.startswith("#"))
    )
    assert body["LSC"][2] == "83127"
    assert body["IRa"][2] == body["IRb"][2] == "26214"
    assert body["SSC"][2] == "18169"
    assert "GC-skew (G−C)/(G+C)" in lines  # the convention is in the header, not a footnote

    report = " ".join((run_dir / "01_structure" / "report.md").read_text("utf-8").split())
    assert "83127" in report and "18169" in report
    assert "LSC → IRa → SSC → IRb" in report


def test_unknown_module_exits_1(tmp_path: Path, data_dir: Path) -> None:
    assert _run(tmp_path, data_dir / "NC_050833.1.gb", extra=["--module", "nope"]) in (1, 2)


def test_word_export_is_reproducible_and_carries_the_golden_numbers(
    tmp_path: Path, data_dir: Path
) -> None:
    """P2-09's Done-when end-to-end: regenerate the document, get an identical file."""
    import hashlib
    import zipfile

    if importlib.util.find_spec("docx") is None:
        pytest.skip("the 'docx' extra is not installed — pip install 'plastid-suite[docx]'")

    assert _run(tmp_path, data_dir / "NC_050833.1.gb", module="gene_table", extra=["--docx"]) == 0
    run_dir = _only_run_dir(tmp_path)
    document = run_dir / "01_gene_table" / "report.docx"
    assert document.is_file()
    first = hashlib.sha256(document.read_bytes()).hexdigest()

    # the manifest indexes it like every other output, with its own checksum
    manifest = json.loads((run_dir / "manifest.json").read_text(encoding="utf-8"))
    entry = next(item for item in manifest["files"] if item["path"].endswith("report.docx"))
    assert entry["sha256"] == first and entry["bytes"] > 0

    # a second run is served from the cache and must restore the very same bytes
    before = _run_dirs(tmp_path)
    assert _run(tmp_path, data_dir / "NC_050833.1.gb", module="gene_table", extra=["--docx"]) == 0
    again = (_run_dirs(tmp_path) - before).pop()
    assert (again / "manifest.json").read_text() != ""
    cached = again / "01_gene_table" / "report.docx"
    assert hashlib.sha256(cached.read_bytes()).hexdigest() == first

    from docx import Document

    with zipfile.ZipFile(document) as archive:
        assert archive.testzip() is None
    opened = Document(str(document))
    text = [paragraph.text for paragraph in opened.paragraphs]
    text += [cell.text for table in opened.tables for row in table.rows for cell in row.cells]
    joined = "\n".join(text)
    assert "Publication gene table" in joined
    assert "132" in joined and "86" in joined and "38" in joined and "8" in joined
    assert "agrees" in joined  # the reconciliation travelled with the document
    assert "Draft — needs human review." in joined  # never mistaken for a checked result


def test_one_run_carries_several_modules_and_combines_their_documents(
    tmp_path: Path, data_dir: Path
) -> None:
    """The run — not the module — is the deliverable: one command, one report, all tables."""
    if importlib.util.find_spec("docx") is None:
        pytest.skip("the 'docx' extra is not installed — pip install 'plastid-suite[docx]'")

    modules = ["genome_stats", "structure", "gene_table"]
    assert (
        _run(
            tmp_path,
            data_dir / "NC_050833.1.gb",
            module=None,
            extra=[part for name in modules for part in ("--module", name)] + ["--docx"],
        )
        == 0
    )
    run_dir = _only_run_dir(tmp_path)
    manifest = json.loads((run_dir / "manifest.json").read_text(encoding="utf-8"))
    validate(manifest)
    assert [entry["id"] for entry in manifest["modules"]] == modules
    assert [entry["status"] for entry in manifest["modules"]] == ["ok"] * 3

    # each module keeps its own numbered directory and its own outputs
    for order, module in enumerate(modules, start=1):
        assert (run_dir / f"{order:02d}_{module}" / "report.md").is_file()
        assert (run_dir / f"{order:02d}_{module}" / "report.docx").is_file()

    combined = run_dir / "report" / "report.docx"
    assert combined.is_file()
    assert any(entry["path"] == "report/report.docx" for entry in manifest["files"])

    from docx import Document

    opened = Document(str(combined))
    text = [paragraph.text for paragraph in opened.paragraphs]
    text += [cell.text for table in opened.tables for row in table.rows for cell in row.cells]
    joined = "\n".join(text)
    for probe in ("Genome statistics", "Genome structure", "Publication gene table"):
        assert probe in joined
    assert "153724" in joined and "83127" in joined and "132" in joined
    # every result table reached the page, and each one got a caption (the document also
    # carries the provenance grid and any pipe table the prose itself contains)
    captioned = [line for line in text if line.startswith("Table ")]
    result_tables = sum(
        len(json.loads((run_dir / f"{order:02d}_{module}" / "result.json").read_text())["tables"])
        for order, module in enumerate(modules, start=1)
    )
    assert len(captioned) == result_tables == 9
    assert len(opened.tables) >= result_tables

    # a fully cached second run reproduces the combined document byte for byte
    first = combined.read_bytes()
    before = _run_dirs(tmp_path)
    assert (
        _run(
            tmp_path,
            data_dir / "NC_050833.1.gb",
            module=None,
            extra=[part for name in modules for part in ("--module", name)] + ["--docx"],
        )
        == 0
    )
    again = (_run_dirs(tmp_path) - before).pop()
    assert (again / "report" / "report.docx").read_bytes() == first


def test_a_run_writes_its_figure_and_a_second_run_writes_the_same_bytes(
    tmp_path: Path, data_dir: Path, tools: dict[str, str | None]
) -> None:
    """P5-01's Done-when, end to end: the figure lands in the run directory and regenerates."""
    import hashlib

    require_tools(tools, "mafft")
    if importlib.util.find_spec("matplotlib") is None:
        pytest.skip("the 'figures' extra is not installed — pip install 'plastid-suite[figures]'")

    def run() -> int:
        return main(
            [
                "run",
                "--module",
                "diversity",
                "--input",
                str(data_dir / "NC_050833.1.gb"),
                "--input",
                str(data_dir / "NC_024936.1.gb"),
                "--out",
                str(tmp_path),
                "--dataset",
                "golden_figures",
                "--set",
                "ir_source=sequence",
                "--figures",
            ]
        )

    assert run() == 0
    first_dir = _only_run_dir(tmp_path, "golden_figures")
    figure = first_dir / "figures" / "pi_window.svg"
    assert figure.is_file(), "the run did not write its figure"
    digest = hashlib.sha256(figure.read_bytes()).hexdigest()

    manifest = json.loads((first_dir / "manifest.json").read_text(encoding="utf-8"))
    entry = next(item for item in manifest["files"] if item["path"] == "figures/pi_window.svg")
    assert entry["module"] == "figures" and entry["sha256"] == digest
    assert manifest["figures"]["rendered"] == [
        "pi_window.svg",
        "pi_window.pdf",
        "pi_window.eps",
    ]
    assert "--figures" in manifest["command"]

    # The second run is served from the cache — and must still get its figure, from the result
    # the first run stored rather than from the module.
    before = _run_dirs(tmp_path, "golden_figures")
    assert run() == 0
    second_dir = (_run_dirs(tmp_path, "golden_figures") - before).pop()
    assert hashlib.sha256((second_dir / "figures" / "pi_window.svg").read_bytes()).hexdigest() == (
        digest
    )

    # The report of that run carries no number its own tables do not (P5-05, on a fresh run).
    from plastid.report.audit import audit_run

    audits = audit_run(first_dir)
    assert audits and all(report.audit.ok for report in audits)


def test_asking_for_a_figure_without_the_extra_is_a_named_exit_3(
    tmp_path: Path, data_dir: Path, tools: dict[str, str | None], monkeypatch: pytest.MonkeyPatch
) -> None:
    """A minimum install runs every pure-Python analysis; asking for a figure says what to add."""
    require_tools(tools, "mafft")
    real_import = importlib.import_module

    def refuse(name: str, package: str | None = None):  # type: ignore[no-untyped-def]
        if name.startswith("matplotlib"):
            raise ImportError("no matplotlib here")
        return real_import(name, package)

    monkeypatch.setattr("plastid.core.deps.import_module", refuse)
    code = main(
        [
            "run",
            "--module",
            "diversity",
            "--input",
            str(data_dir / "NC_050833.1.gb"),
            "--input",
            str(data_dir / "NC_024936.1.gb"),
            "--out",
            str(tmp_path),
            "--dataset",
            "no_figures_extra",
            "--set",
            "ir_source=sequence",
            "--figures",
        ]
    )
    assert code == 3  # the missing extra is a named failure, never a traceback
    assert not (_only_run_dir(tmp_path, "no_figures_extra") / "figures").exists()


def test_naming_a_module_twice_is_a_usage_error(tmp_path: Path, data_dir: Path) -> None:
    """Two runs of one module in a run would write one directory twice; refuse instead."""
    assert (
        _run(
            tmp_path,
            data_dir / "NC_050833.1.gb",
            module=None,
            extra=["--module", "structure", "--module", "structure"],
        )
        == 2
    )
